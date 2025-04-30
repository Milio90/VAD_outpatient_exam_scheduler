import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import pandas as pd
import os
import json
from datetime import datetime
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from tkcalendar import DateEntry  # Import DateEntry for date selection

class PrescriptionSystem:
    def __init__(self, root):
        self.root = root
        self.root.title("Σύστημα Συνταγογράφησης Εξετάσεων")
        self.root.geometry("1200x800")
        
        # Data storage
        self.patient_data = None
        self.selected_patients = []
        self.current_selected_patient = None
        self.examination_selections = {}
        
        # Common examinations
        self.common_blood_tests = [
            "Γενική αίματος", "Na", "K", "ουρία", "κρεατινίνη", "PT",
            "SGOT", "SGPT", "γGT", "LDH", "χολερυθρίνη (ολ)", "Συγκόλληση Αιμοπεταλίων"
        ]
        
        self.additional_tests = [
            "CRP", "FT3", "FT4", "TSH"
        ]
        
        # Separate category for ultrasound
        self.imaging_tests = ["Υπέρηχος καρδιάς"]
        
        # Custom exams - load from file if exists
        self.custom_exams = self.load_custom_exams()
        
        # Exam groups - load from file if exists
        self.exam_groups = self.load_exam_groups()
        
        # Blood test checkboxes
        self.blood_test_vars = {}
        for test in self.common_blood_tests:
            self.blood_test_vars[test] = tk.BooleanVar()
            
        # Additional test checkboxes
        self.additional_test_vars = {}
        for test in self.additional_tests:
            self.additional_test_vars[test] = tk.BooleanVar()
        
        # Imaging test checkboxes
        self.imaging_test_vars = {}
        for test in self.imaging_tests:
            self.imaging_test_vars[test] = tk.BooleanVar()
            
        # Custom test checkboxes
        self.custom_test_vars = {}
        for test in self.custom_exams:
            self.custom_test_vars[test] = tk.BooleanVar()
        
        # Load last used file path if available
        self.last_file_path = self.load_last_file_path()
        
        # Setup UI
        self.setup_ui()
        
        # Auto-load last file if path is saved
        if self.last_file_path and os.path.exists(self.last_file_path):
            self.load_file(self.last_file_path)
    
    def load_last_file_path(self):
        """Load the last used file path from settings"""
        try:
            if os.path.exists("settings.json"):
                with open("settings.json", "r", encoding="utf-8") as f:
                    settings = json.load(f)
                    return settings.get("last_file_path", "")
            return ""
        except Exception as e:
            print(f"Error loading settings: {e}")
            return ""
    
    def save_last_file_path(self, file_path):
        """Save the last used file path to settings"""
        try:
            settings = {}
            if os.path.exists("settings.json"):
                with open("settings.json", "r", encoding="utf-8") as f:
                    settings = json.load(f)
            
            settings["last_file_path"] = file_path
            
            with open("settings.json", "w", encoding="utf-8") as f:
                json.dump(settings, f, ensure_ascii=False)
        except Exception as e:
            print(f"Error saving settings: {e}")
    
    def load_custom_exams(self):
        """Load custom exams from file"""
        try:
            if os.path.exists("custom_exams.json"):
                with open("custom_exams.json", "r", encoding="utf-8") as f:
                    return json.load(f)
            return []
        except Exception as e:
            print(f"Error loading custom exams: {e}")
            return []
    
    def save_custom_exams(self):
        """Save custom exams to file"""
        try:
            with open("custom_exams.json", "w", encoding="utf-8") as f:
                json.dump(self.custom_exams, f, ensure_ascii=False)
        except Exception as e:
            print(f"Error saving custom exams: {e}")
    
    def load_exam_groups(self):
        """Load exam groups from file"""
        try:
            if os.path.exists("exam_groups.json"):
                with open("exam_groups.json", "r", encoding="utf-8") as f:
                    return json.load(f)
            return {}
        except Exception as e:
            print(f"Error loading exam groups: {e}")
            return {}
    
    def save_exam_groups(self):
        """Save exam groups to file"""
        try:
            with open("exam_groups.json", "w", encoding="utf-8") as f:
                json.dump(self.exam_groups, f, ensure_ascii=False)
        except Exception as e:
            print(f"Error saving exam groups: {e}")
    
    def format_amka(self, amka_value):
        """Format AMKA to have exactly 11 digits with preceding zeros"""
        if not amka_value:
            return ""
            
        # Remove any non-digit characters
        amka_digits = re.sub(r'\D', '', str(amka_value))
        
        # Ensure it has exactly 11 digits by adding leading zeros
        return amka_digits.zfill(11)
    
    def setup_ui(self):
        # Main layout with frames
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # File selection area
        file_frame = ttk.LabelFrame(main_frame, text="Επιλογή Αρχείου")
        file_frame.pack(fill=tk.X, padx=5, pady=5)
        
        file_label = ttk.Label(file_frame, text="Επιλογή αρχείου Excel:")
        file_label.pack(side=tk.LEFT, padx=5, pady=5)
        
        self.file_path_label = ttk.Label(file_frame, text="Δεν έχει επιλεγεί αρχείο")
        self.file_path_label.pack(side=tk.LEFT, padx=5, pady=5, expand=True, fill=tk.X)
        
        file_btn = ttk.Button(file_frame, text="Αναζήτηση...", command=self.select_file)
        file_btn.pack(side=tk.RIGHT, padx=5, pady=5)
        
        # Date selection area
        date_frame = ttk.LabelFrame(main_frame, text="Ημερομηνία")
        date_frame.pack(fill=tk.X, padx=5, pady=5)
        
        date_label = ttk.Label(date_frame, text="Επιλογή ημερομηνίας:")
        date_label.pack(side=tk.LEFT, padx=5, pady=5)
        
        # Initialize with current date
        self.date_picker = DateEntry(date_frame, width=12, background='darkblue',
                                    foreground='white', borderwidth=2, date_pattern='dd/MM/yyyy')
        self.date_picker.pack(side=tk.LEFT, padx=5, pady=5)
        
        # Create main content area with patient list and examination selection
        content_frame = ttk.Frame(main_frame)
        content_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Split into left and right panes
        content_frame.columnconfigure(0, weight=2)
        content_frame.columnconfigure(1, weight=3)
        content_frame.rowconfigure(0, weight=1)
        
        # Patient list (left)
        self.setup_patient_list(content_frame)
        
        # Examination selection (right)
        self.setup_examination_frame(content_frame)
        
        # Action buttons at bottom
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X, padx=5, pady=5)
        
        self.apply_btn = ttk.Button(button_frame, text="Εφαρμογή Επιλογών", command=self.apply_selections)
        self.apply_btn.pack(side=tk.LEFT, padx=5, pady=5)
        
        # Add button for exam groups
        groups_btn = ttk.Button(button_frame, text="Διαχείριση Ομάδων Εξετάσεων", command=self.manage_exam_groups)
        groups_btn.pack(side=tk.LEFT, padx=5, pady=5)
        
        self.generate_btn = ttk.Button(button_frame, text="Δημιουργία Εγγράφου", command=self.generate_document)
        self.generate_btn.pack(side=tk.RIGHT, padx=5, pady=5)
        
        # Initial UI state
        self.update_ui_state()
    
    def setup_patient_list(self, parent):
        patient_frame = ttk.LabelFrame(parent, text="Λίστα Ασθενών")
        patient_frame.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        
        # Create a scrollable treeview for patient list
        patient_frame.rowconfigure(0, weight=1)
        patient_frame.columnconfigure(0, weight=1)
        
        # Create the patient table with scrollbar
        table_frame = ttk.Frame(patient_frame)
        table_frame.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        
        # Scrollbars
        y_scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL)
        y_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        x_scrollbar = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL)
        x_scrollbar.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Treeview
        self.patient_tree = ttk.Treeview(
            table_frame, 
            columns=("AM", "Επώνυμο", "Όνομα", "AMKA", "ΚΕΠΑ", "Επιλογή"),
            show='headings',
            yscrollcommand=y_scrollbar.set,
            xscrollcommand=x_scrollbar.set
        )
        
        # Configure scrollbars
        y_scrollbar.config(command=self.patient_tree.yview)
        x_scrollbar.config(command=self.patient_tree.xview)
        
        # Setup columns
        self.patient_tree.heading("AM", text="A.M.")
        self.patient_tree.heading("Επώνυμο", text="Επώνυμο")
        self.patient_tree.heading("Όνομα", text="Όνομα")
        self.patient_tree.heading("AMKA", text="ΑΜΚΑ")
        self.patient_tree.heading("ΚΕΠΑ", text="Απόφαση ΚΕΠΑ")
        self.patient_tree.heading("Επιλογή", text="Επιλογή")
        
        # Column widths
        self.patient_tree.column("AM", width=80, minwidth=50)
        self.patient_tree.column("Επώνυμο", width=120, minwidth=80)
        self.patient_tree.column("Όνομα", width=120, minwidth=80)
        self.patient_tree.column("AMKA", width=120, minwidth=100)
        self.patient_tree.column("ΚΕΠΑ", width=120, minwidth=80)
        self.patient_tree.column("Επιλογή", width=60, minwidth=50)
        
        self.patient_tree.pack(fill=tk.BOTH, expand=True)
        
        # Bind selection event
        self.patient_tree.bind("<<TreeviewSelect>>", self.on_patient_select)
    
    def setup_examination_frame(self, parent):
        exam_frame = ttk.LabelFrame(parent, text="Επιλογή Εξετάσεων")
        exam_frame.grid(row=0, column=1, sticky="nsew", padx=5, pady=5)
        
        # Create a canvas with scrollbar to handle overflow
        canvas = tk.Canvas(exam_frame)
        scrollbar = ttk.Scrollbar(exam_frame, orient="vertical", command=canvas.yview)
        scrollable_frame = ttk.Frame(canvas)
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # Exam groups combobox
        groups_frame = ttk.Frame(scrollable_frame)
        groups_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        groups_label = ttk.Label(groups_frame, text="Ομάδες Εξετάσεων:")
        groups_label.pack(side=tk.LEFT, padx=5, pady=5)
        
        self.group_var = tk.StringVar()
        self.group_combobox = ttk.Combobox(groups_frame, textvariable=self.group_var, width=30)
        self.group_combobox.pack(side=tk.LEFT, padx=5, pady=5)
        self.update_group_combobox()
        
        apply_group_btn = ttk.Button(
            groups_frame, 
            text="Εφαρμογή Ομάδας", 
            command=self.apply_exam_group
        )
        apply_group_btn.pack(side=tk.LEFT, padx=5, pady=5)
        
        # Blood tests groupbox
        blood_frame = ttk.LabelFrame(scrollable_frame, text="Συνήθεις Αιματολογικές Εξετάσεις")
        blood_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Select all blood tests checkbox
        self.select_all_blood = tk.BooleanVar()
        select_all_blood_cb = ttk.Checkbutton(
            blood_frame, 
            text="Επιλογή όλων", 
            variable=self.select_all_blood,
            command=self.toggle_all_blood_tests
        )
        select_all_blood_cb.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for blood tests
        row, col = 1, 0
        for test in self.common_blood_tests:
            cb = ttk.Checkbutton(blood_frame, text=test, variable=self.blood_test_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Additional tests groupbox
        additional_frame = ttk.LabelFrame(scrollable_frame, text="Πρόσθετες Εξετάσεις")
        additional_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Select all additional tests checkbox
        self.select_all_additional = tk.BooleanVar()
        select_all_additional_cb = ttk.Checkbutton(
            additional_frame, 
            text="Επιλογή όλων", 
            variable=self.select_all_additional,
            command=self.toggle_all_additional_tests
        )
        select_all_additional_cb.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for additional tests
        row, col = 1, 0
        for test in self.additional_tests:
            cb = ttk.Checkbutton(additional_frame, text=test, variable=self.additional_test_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Imaging tests groupbox
        imaging_frame = ttk.LabelFrame(scrollable_frame, text="Απεικονιστικές Εξετάσεις")
        imaging_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Select all imaging tests checkbox
        self.select_all_imaging = tk.BooleanVar()
        select_all_imaging_cb = ttk.Checkbutton(
            imaging_frame, 
            text="Επιλογή όλων", 
            variable=self.select_all_imaging,
            command=self.toggle_all_imaging_tests
        )
        select_all_imaging_cb.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for imaging tests
        row, col = 1, 0
        for test in self.imaging_tests:
            cb = ttk.Checkbutton(imaging_frame, text=test, variable=self.imaging_test_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Custom exams groupbox
        self.custom_frame = ttk.LabelFrame(scrollable_frame, text="Προσαρμοσμένες Εξετάσεις")
        self.custom_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Button to add custom exam
        add_custom_btn = ttk.Button(
            self.custom_frame, 
            text="Προσθήκη νέας εξέτασης", 
            command=self.add_custom_exam
        )
        add_custom_btn.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for custom exams
        self.update_custom_exams_ui()
        
        # Notes area
        notes_frame = ttk.LabelFrame(scrollable_frame, text="Παρατηρήσεις")
        notes_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        self.notes_field = tk.Text(notes_frame, height=4, width=40)
        self.notes_field.pack(fill=tk.X, expand=True, padx=5, pady=5)
    
    def toggle_all_blood_tests(self):
        """Toggle all blood tests based on select all checkbox"""
        status = self.select_all_blood.get()
        for var in self.blood_test_vars.values():
            var.set(status)
    
    def toggle_all_additional_tests(self):
        """Toggle all additional tests based on select all checkbox"""
        status = self.select_all_additional.get()
        for var in self.additional_test_vars.values():
            var.set(status)
    
    def toggle_all_imaging_tests(self):
        """Toggle all imaging tests based on select all checkbox"""
        status = self.select_all_imaging.get()
        for var in self.imaging_test_vars.values():
            var.set(status)
    
    def update_custom_exams_ui(self):
        """Update the UI with current custom exams"""
        # Clear existing checkboxes (from row 1 onwards)
        for widget in self.custom_frame.grid_slaves():
            if int(widget.grid_info()["row"]) > 0:
                widget.destroy()
        
        # Create a grid of checkboxes for custom exams
        row, col = 1, 0
        for test in self.custom_exams:
            # Create checkbutton variable if it doesn't exist
            if test not in self.custom_test_vars:
                self.custom_test_vars[test] = tk.BooleanVar()
            
            # Create checkbox frame with delete button
            test_frame = ttk.Frame(self.custom_frame)
            test_frame.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            # Add checkbutton
            cb = ttk.Checkbutton(test_frame, text=test, variable=self.custom_test_vars[test])
            cb.pack(side=tk.LEFT)
            
            # Add delete button
            del_btn = ttk.Button(
                test_frame, 
                text="X", 
                width=2,
                command=lambda t=test: self.delete_custom_exam(t)
            )
            del_btn.pack(side=tk.RIGHT)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
    
    def add_custom_exam(self):
        """Add a new custom exam"""
        new_exam = simpledialog.askstring(
            "Προσθήκη Εξέτασης", 
            "Εισάγετε το όνομα της νέας εξέτασης:"
        )
        
        if new_exam and new_exam.strip():
            new_exam = new_exam.strip()
            
            # Check if it already exists
            if (new_exam in self.common_blood_tests or 
                new_exam in self.additional_tests or 
                new_exam in self.imaging_tests or 
                new_exam in self.custom_exams):
                messagebox.showwarning(
                    "Διπλή Εγγραφή", 
                    "Η εξέταση υπάρχει ήδη!"
                )
                return
            
            # Add to list and update UI
            self.custom_exams.append(new_exam)
            self.custom_test_vars[new_exam] = tk.BooleanVar()
            
            # Save to file
            self.save_custom_exams()
            
            # Update UI
            self.update_custom_exams_ui()
    
    def delete_custom_exam(self, exam_name):
        """Delete a custom exam"""
        # Confirm deletion
        if messagebox.askyesno(
            "Διαγραφή Εξέτασης", 
            f"Είστε σίγουροι ότι θέλετε να διαγράψετε την εξέταση '{exam_name}';"
        ):
            # Remove from list
            if exam_name in self.custom_exams:
                self.custom_exams.remove(exam_name)
            
            # Remove from variables
            if exam_name in self.custom_test_vars:
                del self.custom_test_vars[exam_name]
            
            # Update UI
            self.update_custom_exams_ui()
            
            # Save to file
            self.save_custom_exams()
            
            # Update selections for all patients to remove this exam
            for patient_id in self.examination_selections:
                if "custom_tests" in self.examination_selections[patient_id]:
                    if exam_name in self.examination_selections[patient_id]["custom_tests"]:
                        del self.examination_selections[patient_id]["custom_tests"][exam_name]
    
    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="Επιλογή αρχείου Excel",
            filetypes=[("Excel Files", "*.xlsx *.xls")]
        )
        
        if file_path:
            self.load_file(file_path)
            
            # Ask if user wants to remember this file location
            if messagebox.askyesno(
                "Αποθήκευση τοποθεσίας",
                "Θέλετε το σύστημα να θυμάται αυτό το αρχείο και να το φορτώνει αυτόματα την επόμενη φορά;",
            ):
                self.save_last_file_path(file_path)
    
    def load_file(self, file_path):
        """Load patient data from the specified file path"""
        try:
            self.patient_data = pd.read_excel(file_path)
            
            # Validate data structure
            required_columns = 5  # Hospital ID, Last Name, First Name, AMKA, KEPA
            if self.patient_data.shape[1] < required_columns:
                messagebox.showwarning(
                    "Σφάλμα Μορφής", 
                    "Το αρχείο Excel πρέπει να περιέχει τουλάχιστον 5 στήλες: A.M., Επώνυμο, Όνομα, ΑΜΚΑ, Απόφαση ΚΕΠΑ"
                )
                self.patient_data = None
            else:
                # Set column names if needed
                if len(self.patient_data.columns) >= 5:
                    self.patient_data.columns = ["AM", "Επώνυμο", "Όνομα", "ΑΜΚΑ", "Απόφαση ΚΕΠΑ"] + list(self.patient_data.columns[5:])
                
                self.file_path_label.config(text=os.path.basename(file_path))
                self.load_patients()
        except Exception as e:
            messagebox.showerror("Σφάλμα", f"Σφάλμα φόρτωσης αρχείου: {str(e)}")
            self.patient_data = None
            
        self.update_ui_state()
    
    def load_patients(self):
        if self.patient_data is None:
            return
            
        # Clear existing data
        for item in self.patient_tree.get_children():
            self.patient_tree.delete(item)
        
        # Fill table with patient data
        for idx, row in self.patient_data.iterrows():
            # Format values as strings
            am = str(row["AM"])
            surname = str(row["Επώνυμο"])
            name = str(row["Όνομα"]) if "Όνομα" in row else ""
            amka = self.format_amka(row["ΑΜΚΑ"])  # Format AMKA to have 11 digits
            kepa = str(row["Απόφαση ΚΕΠΑ"]) if "Απόφαση ΚΕΠΑ" in row else ""
            
            # Insert data
            self.patient_tree.insert(
                "", "end", 
                values=(am, surname, name, amka, kepa, "□"),
                iid=am  # Use hospital ID as the item ID
            )
    
    def on_patient_select(self, event):
        selected_items = self.patient_tree.selection()
        if not selected_items:
            return
            
        hospital_id = selected_items[0]  # Get the first selected item
        
        # Save current selections if there's a patient being edited
        if self.current_selected_patient:
            self.save_current_selections()
        
        # Set current patient and load their selections
        self.current_selected_patient = hospital_id
        
        # Load the current patient's selections or initialize empty selections
        self.load_exam_selections(hospital_id)
        
        # Update all checkbox states in the tree view
        self.update_all_checkboxes()
        
        self.update_ui_state()
    
    def update_all_checkboxes(self):
        # Update the checkbox appearances for all patients
        for item_id in self.patient_tree.get_children():
            if item_id in self.selected_patients:
                # Show as checked
                patient_values = self.patient_tree.item(item_id, "values")
                if patient_values[-1] != "✓":  # Only update if needed
                    patient_values = patient_values[:-1] + ("✓",)
                    self.patient_tree.item(item_id, values=patient_values)
            else:
                # Show as unchecked
                patient_values = self.patient_tree.item(item_id, "values")
                if patient_values[-1] != "□":  # Only update if needed
                    patient_values = patient_values[:-1] + ("□",)
                    self.patient_tree.item(item_id, values=patient_values)
    
    def save_current_selections(self):
        """Save the current exam selections for the current patient"""
        if not self.current_selected_patient:
            return
            
        # Collect selections
        blood_tests = {test: var.get() for test, var in self.blood_test_vars.items()}
        additional_tests = {test: var.get() for test, var in self.additional_test_vars.items()}
        imaging_tests = {test: var.get() for test, var in self.imaging_test_vars.items()}
        custom_tests = {test: var.get() for test, var in self.custom_test_vars.items()}
        
        # Get notes
        notes = self.notes_field.get("1.0", "end-1c")
        
        # Save selections
        self.examination_selections[self.current_selected_patient] = {
            "blood_tests": blood_tests,
            "additional_tests": additional_tests,
            "imaging_tests": imaging_tests,
            "custom_tests": custom_tests,
            "notes": notes
        }
        
        # Add to selected patients list if there are any tests selected
        has_selected_tests = (
            any(blood_tests.values()) or 
            any(additional_tests.values()) or
            any(imaging_tests.values()) or
            any(custom_tests.values())
        )
        
        if has_selected_tests and self.current_selected_patient not in self.selected_patients:
            self.selected_patients.append(self.current_selected_patient)
        elif not has_selected_tests and self.current_selected_patient in self.selected_patients:
            self.selected_patients.remove(self.current_selected_patient)
    
    def load_exam_selections(self, hospital_id):
        # Clear current selections
        for var in self.blood_test_vars.values():
            var.set(False)
        for var in self.additional_test_vars.values():
            var.set(False)
        for var in self.imaging_test_vars.values():
            var.set(False)
        for var in self.custom_test_vars.values():
            var.set(False)
        
        # Reset select all checkboxes
        self.select_all_blood.set(False)
        self.select_all_additional.set(False)
        self.select_all_imaging.set(False)
        
        # Clear notes field
        self.notes_field.delete("1.0", tk.END)
            
        # If patient has existing selections, load them
        if hospital_id in self.examination_selections:
            exams = self.examination_selections[hospital_id]
            
            # Set blood test checkboxes
            for test, selected in exams.get("blood_tests", {}).items():
                if test in self.blood_test_vars:
                    self.blood_test_vars[test].set(selected)
                    
            # Set additional test checkboxes
            for test, selected in exams.get("additional_tests", {}).items():
                if test in self.additional_test_vars:
                    self.additional_test_vars[test].set(selected)
            
            # Set imaging test checkboxes
            for test, selected in exams.get("imaging_tests", {}).items():
                if test in self.imaging_test_vars:
                    self.imaging_test_vars[test].set(selected)
            
            # Set custom test checkboxes
            for test, selected in exams.get("custom_tests", {}).items():
                if test in self.custom_test_vars:
                    self.custom_test_vars[test].set(selected)
                    
            # Set notes
            if "notes" in exams:
                self.notes_field.insert("1.0", exams["notes"])
    
    def clear_exam_selections(self):
        # Clear all checkboxes
        for var in self.blood_test_vars.values():
            var.set(False)
        for var in self.additional_test_vars.values():
            var.set(False)
        for var in self.imaging_test_vars.values():
            var.set(False)
        for var in self.custom_test_vars.values():
            var.set(False)
        
        # Reset select all checkboxes
        self.select_all_blood.set(False)
        self.select_all_additional.set(False)
        self.select_all_imaging.set(False)
        
        # Clear notes
        self.notes_field.delete("1.0", tk.END)
    
    def apply_selections(self):
        if not self.current_selected_patient:
            messagebox.showwarning("Προειδοποίηση", "Δεν έχει επιλεγεί ασθενής.")
            return
            
        # Save current selections using the helper method
        self.save_current_selections()
        
        # Update all checkbox states in the tree view
        self.update_all_checkboxes()
        
        self.update_ui_state()
        messagebox.showinfo("Ενημέρωση", "Οι επιλογές εξετάσεων αποθηκεύτηκαν επιτυχώς.")
    
    def generate_document(self):
        # Save current selections first (in case they weren't applied)
        if self.current_selected_patient:
            self.save_current_selections()
        
        # Check if we have any selected patients with examinations
        filtered_selections = {}
        for patient_id in self.selected_patients:
            if patient_id in self.examination_selections:
                # Verify that at least one test is selected
                exams = self.examination_selections[patient_id]
                has_tests = (
                    any(exams.get("blood_tests", {}).values()) or 
                    any(exams.get("additional_tests", {}).values()) or
                    any(exams.get("imaging_tests", {}).values()) or
                    any(exams.get("custom_tests", {}).values())
                )
                if has_tests:
                    filtered_selections[patient_id] = exams
        
        if not filtered_selections:
            messagebox.showwarning("Προειδοποίηση", "Δεν υπάρχουν επιλεγμένες εξετάσεις για κανέναν ασθενή.")
            return
        
        # Create a new document
        doc = Document()
        
        # Set up document style
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(11)
        
        # Add title
        title = doc.add_paragraph("Εξωτερικά Ιατρεία VAD")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].bold = True
        
        subtitle = doc.add_paragraph("Προγραμματισμός -- Συνταγογράφηση Εξετάσεων")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].bold = True
        
        # Add selected date
        date_str = self.date_picker.get_date().strftime("%d.%m.%Y")
        date_paragraph = doc.add_paragraph(f"Ημερομηνία: {date_str}")
        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_paragraph.runs[0].bold = True
        
        # Create table
        table = doc.add_table(rows=1, cols=5)  # 5 columns instead of 6
        table.style = 'Table Grid'
        
        # Table header
        header_cells = table.rows[0].cells
        header_cells[0].text = "α/α"
        header_cells[1].text = "A.M."
        header_cells[2].text = "Επώνυμο"
        header_cells[3].text = "ΑΜΚΑ"
        header_cells[4].text = "Παρατηρήσεις"
        
        # Make header bold
        for i in range(5):
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        
        # Add patient data
        counter = 1
        for hospital_id, exams in filtered_selections.items():
            # Find patient data in treeview
            patient_values = self.patient_tree.item(hospital_id, "values")
            if not patient_values:
                continue
                
            # Add patient row (bold)
            table.add_row()
            row_cells = table.rows[-1].cells
            row_cells[0].text = str(counter)
            row_cells[1].text = patient_values[0]  # AM
            row_cells[2].text = patient_values[1]  # Surname
            row_cells[3].text = patient_values[3]  # AMKA
            
            # Add KEPA info to the notes column
            kepa_info = patient_values[4] if len(patient_values) > 4 and patient_values[4] else ""
            row_cells[4].text = f"Απόφαση ΚΕΠΑ: {kepa_info}" if kepa_info else ""
            
            # Make patient row bold
            for i in range(5):
                for paragraph in row_cells[i].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            
            # Add blood tests row
            blood_tests_text = ""
            for test, selected in exams["blood_tests"].items():
                if selected:
                    blood_tests_text += f"{test}, "
            
            # Additional tests
            for test, selected in exams["additional_tests"].items():
                if selected:
                    blood_tests_text += f"{test}, "
                    
            # Custom tests (non-imaging)
            for test, selected in exams.get("custom_tests", {}).items():
                if selected:
                    blood_tests_text += f"{test}, "
            
            if blood_tests_text:
                blood_tests_text = blood_tests_text.rstrip(", ")
                table.add_row()
                row_cells = table.rows[-1].cells
                row_cells[0].text = blood_tests_text
                # Merge cells for lab tests to span across columns
                for i in range(1, 4):
                    row_cells[0].merge(row_cells[i])
                row_cells[4].text = ""
            
            # Imaging tests in separate row
            imaging_tests_text = ""
            for test, selected in exams.get("imaging_tests", {}).items():
                if selected:
                    imaging_tests_text += f"{test}, "
            
            if imaging_tests_text:
                imaging_tests_text = imaging_tests_text.rstrip(", ")
                table.add_row()
                row_cells = table.rows[-1].cells
                row_cells[0].text = imaging_tests_text
                # Merge cells for imaging tests to span across columns
                for i in range(1, 4):
                    row_cells[0].merge(row_cells[i])
                row_cells[4].text = ""
            
            # Notes row if needed
            notes_text = exams.get("notes", "").strip()
            if notes_text:
                table.add_row()
                row_cells = table.rows[-1].cells
                row_cells[0].text = "Σημειώσεις: " + notes_text
                # Merge cells for notes to span across columns
                for i in range(1, 5):
                    row_cells[0].merge(row_cells[i])
            
            counter += 1
        
        # Save document - use the selected date in the filename
        date_for_filename = self.date_picker.get_date().strftime("%d.%m.%Y")
        save_path = filedialog.asksaveasfilename(
            title="Αποθήκευση Εγγράφου",
            initialfile=f"Συνταγογράφηση VAD ({date_for_filename}).docx",
            filetypes=[("Word Documents", "*.docx")]
        )
        
        if save_path:
            if not save_path.endswith('.docx'):
                save_path += '.docx'
            doc.save(save_path)
            messagebox.showinfo("Επιτυχία", f"Το έγγραφο αποθηκεύτηκε επιτυχώς στο: {save_path}")
    
    def update_ui_state(self):
        file_loaded = self.patient_data is not None
        patient_selected = self.current_selected_patient is not None
        
        # Update button states
        self.apply_btn.config(state=tk.NORMAL if patient_selected else tk.DISABLED)
        self.generate_btn.config(state=tk.NORMAL if file_loaded and self.examination_selections else tk.DISABLED)
    
    def update_group_combobox(self):
        """Update the exam groups dropdown"""
        group_names = list(self.exam_groups.keys())
        self.group_combobox['values'] = group_names
        if group_names:
            self.group_combobox.current(0)
    
    def apply_exam_group(self):
        """Apply the selected exam group to the current patient"""
        if not self.current_selected_patient:
            messagebox.showwarning("Προειδοποίηση", "Δεν έχει επιλεγεί ασθενής.")
            return
            
        selected_group = self.group_var.get()
        if not selected_group or selected_group not in self.exam_groups:
            messagebox.showwarning("Προειδοποίηση", "Δεν έχει επιλεγεί έγκυρη ομάδα εξετάσεων.")
            return
            
        # Get the group's exams
        group_exams = self.exam_groups[selected_group]
        
        # Apply selections
        # First clear all current selections
        self.clear_exam_selections()
        
        # Set blood test checkboxes
        for test, selected in group_exams.get("blood_tests", {}).items():
            if test in self.blood_test_vars:
                self.blood_test_vars[test].set(selected)
                
        # Set additional test checkboxes
        for test, selected in group_exams.get("additional_tests", {}).items():
            if test in self.additional_test_vars:
                self.additional_test_vars[test].set(selected)
        
        # Set imaging test checkboxes
        for test, selected in group_exams.get("imaging_tests", {}).items():
            if test in self.imaging_test_vars:
                self.imaging_test_vars[test].set(selected)
        
        # Set custom test checkboxes
        for test, selected in group_exams.get("custom_tests", {}).items():
            if test in self.custom_test_vars:
                self.custom_test_vars[test].set(selected)
        
        messagebox.showinfo("Ενημέρωση", f"Η ομάδα '{selected_group}' εφαρμόστηκε επιτυχώς.")
    
    def manage_exam_groups(self):
        """Open dialog to manage exam groups"""
        groups_window = tk.Toplevel(self.root)
        groups_window.title("Διαχείριση Ομάδων Εξετάσεων")
        groups_window.geometry("600x500")
        groups_window.transient(self.root)
        groups_window.grab_set()
        
        # Main frame
        main_frame = ttk.Frame(groups_window, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Groups list frame
        groups_frame = ttk.LabelFrame(main_frame, text="Διαθέσιμες Ομάδες")
        groups_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Create a listbox with scrollbar for groups
        groups_list_frame = ttk.Frame(groups_frame)
        groups_list_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        scrollbar = ttk.Scrollbar(groups_list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.groups_listbox = tk.Listbox(groups_list_frame, yscrollcommand=scrollbar.set)
        self.groups_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        scrollbar.config(command=self.groups_listbox.yview)
        
        # Populate listbox
        for group_name in self.exam_groups.keys():
            self.groups_listbox.insert(tk.END, group_name)
        
        # Action buttons
        buttons_frame = ttk.Frame(main_frame)
        buttons_frame.pack(fill=tk.X, padx=5, pady=5)
        
        create_btn = ttk.Button(
            buttons_frame, 
            text="Δημιουργία Νέας Ομάδας",
            command=lambda: self.create_exam_group(groups_window)
        )
        create_btn.pack(side=tk.LEFT, padx=5, pady=5)
        
        edit_btn = ttk.Button(
            buttons_frame, 
            text="Επεξεργασία Ομάδας",
            command=lambda: self.edit_exam_group(groups_window)
        )
        edit_btn.pack(side=tk.LEFT, padx=5, pady=5)
        
        delete_btn = ttk.Button(
            buttons_frame, 
            text="Διαγραφή Ομάδας",
            command=lambda: self.delete_exam_group(groups_window)
        )
        delete_btn.pack(side=tk.LEFT, padx=5, pady=5)
        
        # Close button
        close_frame = ttk.Frame(main_frame)
        close_frame.pack(fill=tk.X, padx=5, pady=5)
        
        close_btn = ttk.Button(
            close_frame,
            text="Κλείσιμο",
            command=groups_window.destroy
        )
        close_btn.pack(side=tk.RIGHT, padx=5, pady=5)
    
    def create_exam_group(self, parent_window):
        """Create a new exam group"""
        # First save current selections if a patient is selected
        if self.current_selected_patient:
            self.save_current_selections()
        
        # Ask for group name
        group_name = simpledialog.askstring(
            "Νέα Ομάδα Εξετάσεων",
            "Εισάγετε όνομα για τη νέα ομάδα:",
            parent=parent_window
        )
        
        if not group_name or not group_name.strip():
            return
            
        group_name = group_name.strip()
        
        # Check if name already exists
        if group_name in self.exam_groups:
            messagebox.showwarning(
                "Διπλή Εγγραφή",
                "Υπάρχει ήδη ομάδα με αυτό το όνομα!"
            )
            return
        
        # Create group window
        group_window = tk.Toplevel(parent_window)
        group_window.title(f"Δημιουργία Ομάδας: {group_name}")
        group_window.geometry("700x600")
        group_window.transient(parent_window)
        group_window.grab_set()
        
        self.setup_group_edit_ui(group_window, group_name)
    
    def edit_exam_group(self, parent_window):
        """Edit an existing exam group"""
        # Get selected group
        selection = self.groups_listbox.curselection()
        if not selection:
            messagebox.showwarning(
                "Προειδοποίηση",
                "Δεν έχει επιλεγεί ομάδα για επεξεργασία!"
            )
            return
            
        group_name = self.groups_listbox.get(selection[0])
        
        # Create edit window
        group_window = tk.Toplevel(parent_window)
        group_window.title(f"Επεξεργασία Ομάδας: {group_name}")
        group_window.geometry("700x600")
        group_window.transient(parent_window)
        group_window.grab_set()
        
        self.setup_group_edit_ui(group_window, group_name, editing=True)
    
    def setup_group_edit_ui(self, window, group_name, editing=False):
        """Set up UI for creating/editing an exam group"""
        # Blood test variables
        blood_vars = {}
        for test in self.common_blood_tests:
            blood_vars[test] = tk.BooleanVar()
            
        # Additional test variables
        additional_vars = {}
        for test in self.additional_tests:
            additional_vars[test] = tk.BooleanVar()
        
        # Imaging test variables
        imaging_vars = {}
        for test in self.imaging_tests:
            imaging_vars[test] = tk.BooleanVar()
            
        # Custom test variables
        custom_vars = {}
        for test in self.custom_exams:
            custom_vars[test] = tk.BooleanVar()
        
        # If editing, load existing selections
        if editing and group_name in self.exam_groups:
            group_data = self.exam_groups[group_name]
            
            # Set blood test checkboxes
            for test, selected in group_data.get("blood_tests", {}).items():
                if test in blood_vars:
                    blood_vars[test].set(selected)
                    
            # Set additional test checkboxes
            for test, selected in group_data.get("additional_tests", {}).items():
                if test in additional_vars:
                    additional_vars[test].set(selected)
            
            # Set imaging test checkboxes
            for test, selected in group_data.get("imaging_tests", {}).items():
                if test in imaging_vars:
                    imaging_vars[test].set(selected)
            
            # Set custom test checkboxes
            for test, selected in group_data.get("custom_tests", {}).items():
                if test in custom_vars:
                    custom_vars[test].set(selected)
        
        # Main frame with scrolling
        main_frame = ttk.Frame(window, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Create a canvas with scrollbar to handle overflow
        canvas = tk.Canvas(main_frame)
        scrollbar = ttk.Scrollbar(main_frame, orient="vertical", command=canvas.yview)
        scrollable_frame = ttk.Frame(canvas)
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # Blood tests groupbox
        blood_frame = ttk.LabelFrame(scrollable_frame, text="Συνήθεις Αιματολογικές Εξετάσεις")
        blood_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Select all blood tests checkbox
        select_all_blood = tk.BooleanVar()
        select_all_blood_cb = ttk.Checkbutton(
            blood_frame, 
            text="Επιλογή όλων", 
            variable=select_all_blood,
            command=lambda: self.toggle_vars(blood_vars, select_all_blood.get())
        )
        select_all_blood_cb.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for blood tests
        row, col = 1, 0
        for test in self.common_blood_tests:
            cb = ttk.Checkbutton(blood_frame, text=test, variable=blood_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Additional tests groupbox
        additional_frame = ttk.LabelFrame(scrollable_frame, text="Πρόσθετες Εξετάσεις")
        additional_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Select all additional tests checkbox
        select_all_additional = tk.BooleanVar()
        select_all_additional_cb = ttk.Checkbutton(
            additional_frame, 
            text="Επιλογή όλων", 
            variable=select_all_additional,
            command=lambda: self.toggle_vars(additional_vars, select_all_additional.get())
        )
        select_all_additional_cb.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for additional tests
        row, col = 1, 0
        for test in self.additional_tests:
            cb = ttk.Checkbutton(additional_frame, text=test, variable=additional_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Imaging tests groupbox
        imaging_frame = ttk.LabelFrame(scrollable_frame, text="Απεικονιστικές Εξετάσεις")
        imaging_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Select all imaging tests checkbox
        select_all_imaging = tk.BooleanVar()
        select_all_imaging_cb = ttk.Checkbutton(
            imaging_frame, 
            text="Επιλογή όλων", 
            variable=select_all_imaging,
            command=lambda: self.toggle_vars(imaging_vars, select_all_imaging.get())
        )
        select_all_imaging_cb.grid(row=0, column=0, columnspan=3, sticky="w", padx=5, pady=2)
        
        # Create a grid of checkboxes for imaging tests
        row, col = 1, 0
        for test in self.imaging_tests:
            cb = ttk.Checkbutton(imaging_frame, text=test, variable=imaging_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Custom exams groupbox
        custom_frame = ttk.LabelFrame(scrollable_frame, text="Προσαρμοσμένες Εξετάσεις")
        custom_frame.pack(fill=tk.X, expand=True, padx=5, pady=5)
        
        # Create a grid of checkboxes for custom exams
        row, col = 0, 0
        for test in self.custom_exams:
            cb = ttk.Checkbutton(custom_frame, text=test, variable=custom_vars[test])
            cb.grid(row=row, column=col, sticky="w", padx=5, pady=2)
            
            col += 1
            if col > 2:  # 3 columns layout
                col = 0
                row += 1
        
        # Action buttons
        buttons_frame = ttk.Frame(window)
        buttons_frame.pack(fill=tk.X, padx=5, pady=10)
        
        save_btn = ttk.Button(
            buttons_frame,
            text="Αποθήκευση",
            command=lambda: self.save_group(
                window, group_name, blood_vars, additional_vars, 
                imaging_vars, custom_vars
            )
        )
        save_btn.pack(side=tk.RIGHT, padx=5, pady=5)
        
        cancel_btn = ttk.Button(
            buttons_frame,
            text="Άκυρο",
            command=window.destroy
        )
        cancel_btn.pack(side=tk.RIGHT, padx=5, pady=5)
    
    def toggle_vars(self, vars_dict, state):
        """Helper to toggle all variables in a dictionary"""
        for var in vars_dict.values():
            var.set(state)
    
    def save_group(self, window, group_name, blood_vars, additional_vars, imaging_vars, custom_vars):
        """Save the group's exam selections"""
        # Collect selections
        blood_tests = {test: var.get() for test, var in blood_vars.items()}
        additional_tests = {test: var.get() for test, var in additional_vars.items()}
        imaging_tests = {test: var.get() for test, var in imaging_vars.items()}
        custom_tests = {test: var.get() for test, var in custom_vars.items()}
        
        # Check if at least one exam is selected
        has_selections = (
            any(blood_tests.values()) or 
            any(additional_tests.values()) or
            any(imaging_tests.values()) or
            any(custom_tests.values())
        )
        
        if not has_selections:
            messagebox.showwarning(
                "Προειδοποίηση",
                "Πρέπει να επιλέξετε τουλάχιστον μία εξέταση!",
                parent=window
            )
            return
        
        # Save group
        self.exam_groups[group_name] = {
            "blood_tests": blood_tests,
            "additional_tests": additional_tests,
            "imaging_tests": imaging_tests,
            "custom_tests": custom_tests
        }
        
        # Save to file
        self.save_exam_groups()
        
        # Update the groups combobox
        self.update_group_combobox()
        
        # Update the groups listbox if it exists
        if hasattr(self, 'groups_listbox'):
            self.groups_listbox.delete(0, tk.END)
            for name in self.exam_groups.keys():
                self.groups_listbox.insert(tk.END, name)
        
        # Close the window
        window.destroy()
        
        messagebox.showinfo("Ενημέρωση", f"Η ομάδα '{group_name}' αποθηκεύτηκε επιτυχώς.")
    
    def delete_exam_group(self, parent_window):
        """Delete an exam group"""
        # Get selected group
        selection = self.groups_listbox.curselection()
        if not selection:
            messagebox.showwarning(
                "Προειδοποίηση",
                "Δεν έχει επιλεγεί ομάδα για διαγραφή!"
            )
            return
            
        group_name = self.groups_listbox.get(selection[0])
        
        # Confirm deletion
        if messagebox.askyesno(
            "Διαγραφή Ομάδας",
            f"Είστε σίγουροι ότι θέλετε να διαγράψετε την ομάδα '{group_name}';",
            parent=parent_window
        ):
            # Remove from dict
            if group_name in self.exam_groups:
                del self.exam_groups[group_name]
            
            # Save to file
            self.save_exam_groups()
            
            # Update the groups list
            self.groups_listbox.delete(0, tk.END)
            for name in self.exam_groups.keys():
                self.groups_listbox.insert(tk.END, name)
            
            # Update the combobox
            self.update_group_combobox()
            
            messagebox.showinfo(
                "Ενημέρωση", 
                f"Η ομάδα '{group_name}' διαγράφηκε επιτυχώς.",
                parent=parent_window
            )


if __name__ == "__main__":
    # Verify required dependencies
    try:
        import tkcalendar
    except ImportError:
        print("Error: tkcalendar module is not installed.")
        print("Please install it using: pip install tkcalendar")
        import sys
        sys.exit(1)
    
    root = tk.Tk(className='VADprescriptionsystem')
    app = PrescriptionSystem(root)
    root.mainloop()
